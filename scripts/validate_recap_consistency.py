#!/usr/bin/env python3
"""Validate report structure, dates, source audit, and Word/HTML consistency."""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from html.parser import HTMLParser
from pathlib import Path

from docx import Document


for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")


REQUIRED_HEADINGS = [
    "1. Executive Tape Snapshot",
    "2. Rotation & Leadership Dashboard",
    "4. Factor, Microstructure & HF Positioning",
    "6. Technology Sector Scan",
    "10. Index Inclusion & Option Risk",
    "15. Sources",
]
REQUIRED_DATA_FILES = [
    "quotes.json",
    "breadth.json",
    "rotation_engine.json",
    "options_chain.json",
    "index_inclusion_watch.json",
    "daily_movers.json",
    "source_audit.json",
]


class TextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        self.parts.append(data)

    @property
    def text(self) -> str:
        return re.sub(r"\s+", " ", " ".join(self.parts)).strip()


def word_text_and_metrics(path: Path) -> tuple[str, dict]:
    document = Document(str(path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    tables = [row for table in document.tables for row in table.rows]
    with zipfile.ZipFile(path) as archive:
        images = len([x for x in archive.namelist() if x.startswith("word/media/")])
    return " ".join(paragraphs + [cell.text for table in document.tables for row in table.rows for cell in row.cells]), {
        "paragraphs": len(paragraphs),
        "tables": len(document.tables),
        "table_rows": len(tables),
        "images": images,
    }


def locate_source_root(report: Path, report_date: str, explicit: Path | None) -> Path | None:
    candidates = []
    if explicit:
        candidates.append(explicit)
    candidates.extend([report.parent / report_date, report.parent.parent / report_date])
    for candidate in candidates:
        if (candidate / "source_audit.json").is_file():
            return candidate
    return None


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", required=True, type=Path)
    parser.add_argument("--html", required=True, type=Path)
    parser.add_argument("--report-date-bjt", required=True)
    parser.add_argument("--us-trade-date", required=True)
    parser.add_argument("--source-root", type=Path)
    parser.add_argument("--previous-html", type=Path)
    args = parser.parse_args()
    errors: list[str] = []
    warnings: list[str] = []
    if not args.report.is_file() or args.report.stat().st_size == 0:
        errors.append(f"Word missing or empty: {args.report}")
    if not args.html.is_file() or args.html.stat().st_size < 1000:
        errors.append(f"HTML missing or too small: {args.html}")
    if errors:
        print(json.dumps({"ok": False, "errors": errors, "warnings": warnings}, ensure_ascii=False))
        return 1
    word_text, metrics = word_text_and_metrics(args.report)
    html_parser = TextParser()
    html_raw = args.html.read_text(encoding="utf-8", errors="replace")
    html_parser.feed(html_raw)
    html_text = html_parser.text
    for token in [args.report_date_bjt, args.us_trade_date]:
        if token not in word_text:
            errors.append(f"Word does not contain expected date: {token}")
        if token not in html_text:
            errors.append(f"HTML does not contain expected date: {token}")
    for heading in REQUIRED_HEADINGS:
        if heading not in word_text:
            errors.append(f"Missing Word section: {heading}")
        if heading not in html_text:
            errors.append(f"Missing HTML section: {heading}")
    if metrics["tables"] < 5:
        errors.append(f"Too few Word tables: {metrics['tables']}")
    if metrics["images"] < 10:
        errors.append(f"Too few embedded Word images: {metrics['images']}")
    html_tables = len(re.findall(r"<table\b", html_raw, flags=re.IGNORECASE))
    html_images = len(re.findall(r"<img\b", html_raw, flags=re.IGNORECASE))
    if html_tables < metrics["tables"]:
        errors.append(f"HTML table count {html_tables} is below Word table count {metrics['tables']}")
    if html_images < metrics["images"]:
        errors.append(f"HTML image count {html_images} is below Word image count {metrics['images']}")
    percentage_tokens = re.findall(r"(?<!\d)[+-]?\d+(?:\.\d+)?%", word_text)[:30]
    missing_tokens = [token for token in percentage_tokens if token not in html_text]
    if missing_tokens:
        errors.append(f"HTML is missing Word numeric tokens: {missing_tokens[:5]}")
    stale_tokens = [
        r"D:\\codex\\outputs\\us_market_daily_recaps\\",
        r"Path\(r?['\"]D:\\codex\\outputs\\us_market_daily_recaps",
    ]
    for pattern in stale_tokens:
        if re.search(pattern, word_text) or re.search(pattern, html_text):
            errors.append(f"Template/path residue detected: {pattern}")
    source_root = locate_source_root(args.report, args.report_date_bjt, args.source_root)
    source_audit = {}
    if source_root is None:
        errors.append("source_audit.json not found; source capture time cannot be verified")
    else:
        for name in REQUIRED_DATA_FILES:
            path = source_root / name
            if not path.is_file() or path.stat().st_size == 0:
                errors.append(f"Missing or empty source artifact: {path}")
        try:
            source_audit = json.loads((source_root / "source_audit.json").read_text(encoding="utf-8"))
            if not source_audit.get("captured_at_utc"):
                errors.append("source_audit.json has no captured_at_utc")
            if source_audit.get("report_date_bjt") != args.report_date_bjt or source_audit.get("us_trade_date") != args.us_trade_date:
                errors.append("source_audit.json date mapping does not match report")
            if source_audit.get("quality", {}).get("degraded"):
                warnings.extend(source_audit.get("quality", {}).get("warnings", []))
            movers_path = source_root / "daily_movers.json"
            movers = json.loads(movers_path.read_text(encoding="utf-8"))
            if not isinstance(movers.get("selected"), list):
                errors.append("daily_movers.json selected is not a list")
            if movers.get("selection_method", "").startswith("fixed"):
                errors.append("daily_movers.json uses a fixed-list selection method")
            if movers.get("universe_count", 0) == 0:
                warnings.append("daily_movers_universe_empty")
        except (OSError, json.JSONDecodeError) as exc:
            errors.append(f"Invalid source_audit.json: {exc}")
    if args.previous_html and args.previous_html.is_file():
        previous_parser = TextParser()
        previous_parser.feed(args.previous_html.read_text(encoding="utf-8", errors="replace"))
        previous_text = previous_parser.text
        if previous_text == html_text:
            errors.append("Current HTML is byte/text-identical to the previous report")
        else:
            warnings.append(f"previous_html_diff_chars={len(html_text) - len(previous_text)}")
    else:
        warnings.append("previous report unavailable; cross-report diff not computed")
    result = {"ok": not errors, "errors": errors, "warnings": warnings, "metrics": metrics, "source_audit": source_audit}
    print(json.dumps(result, ensure_ascii=False))
    return 0 if not errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
